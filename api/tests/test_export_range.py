"""Bulk download over a date range: Excel and a ZIP of PDFs, for branch admins and the central admin."""

import io
import zipfile

import pytest
from openpyxl import load_workbook

from tests.test_export_stock import API, make  # noqa: F401  (make() creates a SUBMITTED invoice of Badges)

D1, D2, D3 = "2027-03-01", "2027-03-02", "2027-03-05"


@pytest.fixture(scope="module")
def prods(client):
    return {p["name"]: p["id"] for p in client.get(f"{API}/products").json()}


@pytest.fixture(scope="module")
def made(client, people, prods):
    def mk(who, company, day):
        r = client.post(f"{API}/invoices", headers=people[who]["h"], json={
            "company_name": company, "invoice_date": day, "action": "submit", "items": [{"product_id": prods["Caps"], "quantity": 1}]})
        assert r.status_code == 201, r.text
        return r.json()
    inv = {"a1": mk("til", "Range A1", D1), "a2": mk("til", "Range A2", D2), "a3": mk("til", "Range A3", D3), "b1": mk("sbr", "Range B1", D2)}
    draft = client.post(f"{API}/invoices", headers=people["til"]["h"], json={"company_name": "Range Draft", "invoice_date": D2, "action": "draft", "items": []})
    assert draft.status_code == 201
    return inv


def sheets(resp):
    assert resp.status_code == 200, resp.text
    return load_workbook(io.BytesIO(resp.content)).sheetnames


def test_excel_covers_the_whole_range_and_only_issued_invoices(client, people, made):
    adm = people["til_admin"]["h"]
    r = client.get(f"{API}/exports/daily-invoices", headers=adm, params={"date_from": D1, "date_to": D2})
    names = sheets(r)
    assert r.headers["content-disposition"] == f'attachment; filename="invoices-TIL-{D1}_to_{D2}.xlsx"'
    assert len([n for n in names if n != "Summary"]) == 2 and not any("Draft" in n for n in names)        # A1 + A2; no draft; A3 is outside
    one = client.get(f"{API}/exports/daily-invoices", headers=adm, params={"date": D3})
    assert len(sheets(one)) == 2 and one.headers["content-disposition"].endswith(f'invoices-TIL-{D3}.xlsx"')   # a single day still works as before
    wb = load_workbook(io.BytesIO(r.content))
    assert "to" in wb["Summary"]["A1"].value


def test_central_admin_range_all_branches_or_one(client, people, made):
    su = people["super"]["h"]
    allb = sheets(client.get(f"{API}/exports/daily-invoices", headers=su, params={"date_from": D1, "date_to": D3}))
    assert len(allb) - 1 >= 4
    sbr = sheets(client.get(f"{API}/exports/daily-invoices", headers=su, params={"date_from": D1, "date_to": D3, "branch": "SBR"}))
    assert len(sbr) - 1 == 1


def test_zip_of_pdfs(client, people, made):
    adm = people["til_admin"]["h"]
    r = client.get(f"{API}/exports/invoice-pdfs", headers=adm, params={"date_from": D1, "date_to": D2})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert len(z.namelist()) == 2 and all(n.endswith(".pdf") for n in z.namelist())
    assert z.read(z.namelist()[0]).startswith(b"%PDF")
    su = client.get(f"{API}/exports/invoice-pdfs", headers=people["super"]["h"], params={"date_from": D1, "date_to": D3})
    names = zipfile.ZipFile(io.BytesIO(su.content)).namelist()
    assert any(n.startswith("TIL/") for n in names) and any(n.startswith("SBR/") for n in names)           # grouped by branch for ALL
    assert client.get(f"{API}/exports/invoice-pdfs", headers=adm, params={"date_from": "2030-01-01", "date_to": "2030-01-02"}).status_code == 404


def test_range_rules_and_permissions(client, people, made):
    adm, user = people["til_admin"]["h"], people["til"]["h"]
    q = {"date_from": D1, "date_to": D2}
    for path in ("daily-invoices", "invoice-pdfs"):
        url = f"{API}/exports/{path}"
        assert client.get(url, params=q).status_code == 401
        assert client.get(url, headers=user, params=q).status_code == 403                                    # an ordinary user cannot
        assert client.get(url, headers=adm, params={**q, "branch": "SBR"}).status_code == 403                # nor another branch
        assert client.get(url, headers=adm, params={"date_from": D2, "date_to": D1}).status_code == 422       # backwards
        assert client.get(url, headers=adm, params={"date_from": "2026-01-01", "date_to": "2027-12-31"}).status_code == 422   # too long
        assert client.get(url, headers=adm, params={"date_from": D1}).status_code == 422                     # one end missing
