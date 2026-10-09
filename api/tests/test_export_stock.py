"""Daily Excel export and simple stock."""

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db import engine

API = "/api/v1"
DAY = "2036-04-04"


@pytest.fixture(scope="module")
def prods(client):
    return {p["name"]: p["id"] for p in client.get(f"{API}/products").json()}


def make(client, people, prods, who, company, qty, day=DAY, action="submit", **extra):
    r = client.post(f"{API}/invoices", headers=people[who]["h"], json={
        "company_name": company, "invoice_date": day, "action": action,
        "items": [{"product_id": prods["Badges"], "quantity": qty}], **extra})
    assert r.status_code == 201, r.text
    return r.json()


def book(resp):
    assert resp.status_code == 200, resp.text
    return load_workbook(io.BytesIO(resp.content))


@pytest.fixture(scope="module")
def invoices(client, people, prods):
    out = {}
    out["til1"] = make(client, people, prods, "til", "Sahyadri Precision Components Pvt. Ltd.", 10,
                       payments=[{"mode": "UPI", "amount": "50", "reference": "UTR 9"}, {"mode": "CASH", "amount": "7"}])
    out["til2"] = make(client, people, prods, "til", "Bad/Name:With*Chars?[x]", 5)
    out["sbr1"] = make(client, people, prods, "sbr", "SBR Customer", 20)
    make(client, people, prods, "til", "A draft is not issued", 1, action="draft")           # excluded
    make(client, people, prods, "til", "Another day", 1, day="2036-04-05")                     # excluded
    cancelled = make(client, people, prods, "til", "Cancelled one", 1)
    client.post(f"{API}/invoices/{cancelled['id']}/cancel", json={"reason": "duplicate"}, headers=people["til_admin"]["h"])
    return out


# ---------------------------------------------------------------- Excel
def test_branch_admin_gets_only_their_branch_one_tab_per_invoice(client, people, invoices):
    wb = book(client.get(f"{API}/exports/daily-invoices", params={"date": DAY}, headers=people["til_admin"]["h"]))
    names = wb.sheetnames
    assert names[0] == "Summary" and len(names) == 3                                      # summary + the 2 issued TIL invoices
    n1 = invoices["til1"]["invoice_number"].split("-")[2]
    assert any(n.startswith(f"TIL-{n1} Sahyadri Precision") for n in names)               # branch code + number + company
    assert all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in names)             # Excel's tab rules
    assert any("BadNameWithChars" in n for n in names)                                    # forbidden characters removed
    assert not any(n.startswith("SBR") for n in names)                                    # never another branch
    summary = wb["Summary"]
    cells = [c.value for row in summary.iter_rows() for c in row if c.value is not None]
    assert "A draft is not issued" not in cells and "Cancelled one" not in cells and "Another day" not in cells
    assert "SBR Customer" not in cells


def test_invoice_tab_matches_the_invoice_layout(client, people, invoices):
    wb = book(client.get(f"{API}/exports/daily-invoices", params={"date": DAY}, headers=people["til_admin"]["h"]))
    ws = next(wb[n] for n in wb.sheetnames if "Sahyadri" in n)
    text_ = " | ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    for expected in ("PROFORMA INVOICE", "Mahratta Chamber of Commerce", "GSTIN - 27AAATM5559Q1ZS", invoices["til1"]["invoice_number"],
                     "Company Name  : Sahyadri Precision", "Particulars", "HSN Code", "Basic Amount", "CGST AMT",
                     "Water Bottle", "Total …", "Rounded off Amount..", "For MCCIA", "Authorized Signatory",
                     "Payment Details : UPI Rs. 50.00 (UTR 9) + Cash Rs. 7.00"):
        assert expected in text_, expected
    header_row = next(r for r in ws.iter_rows() if r[0].value == "Sr.")
    assert [c.value for c in header_row][:5] == ["Sr.", "Particulars", "HSN Code", "Rate", "Qty."]
    assert header_row[0].fill.fgColor.rgb.endswith("C5D9F1")                                # the blue header
    badges = next(r for r in ws.iter_rows() if r[1].value == "Badges")
    assert badges[4].value == 10 and round(badges[5].value, 2) == 48.0 and round(badges[10].value, 2) == 56.64
    total = next(r for r in ws.iter_rows() if r[1].value == "Total …")
    assert total[4].value == 10 and total[1].fill.fgColor.rgb.endswith("C5D9F1")
    assert sum(1 for r in ws.iter_rows() if r[1].value and r[2].value and str(r[2].value).isdigit()) == 39   # all catalogue rows


def test_central_admin_combined_and_separate_files(client, people, invoices):
    su = people["super"]["h"]
    combined = book(client.get(f"{API}/exports/daily-invoices", params={"date": DAY, "branch": "ALL"}, headers=su))
    assert len(combined.sheetnames) == 4 and combined.sheetnames[0] == "Summary"
    order = [n.split("-")[0] for n in combined.sheetnames[1:]]
    assert order == sorted(order) and set(order) == {"SBR", "TIL"}                         # grouped by branch
    rows = [[c.value for c in r] for r in combined["Summary"].iter_rows(min_row=4)]
    totals = {r[0]: r[6] for r in rows if r[0] and str(r[0]).endswith(" total")}
    assert totals == {"SBR total": 113.0, "TIL total": 85.0}                                   # 20 badges; 10 + 5 badges (payable, rounded)
    flat = " ".join(str(v) for r in rows for v in r if v is not None)
    assert "SBR total" in flat and "TIL total" in flat and "All branches" in flat
    only = book(client.get(f"{API}/exports/daily-invoices", params={"date": DAY, "branch": "SBR"}, headers=su))
    assert len(only.sheetnames) == 2 and only.sheetnames[1].startswith("SBR-")
    empty = book(client.get(f"{API}/exports/daily-invoices", params={"date": DAY, "branch": "AHL"}, headers=su))
    assert empty.sheetnames == ["Summary"]                                                  # a branch with nothing that day
    default = book(client.get(f"{API}/exports/daily-invoices", params={"date": DAY}, headers=su))
    assert len(default.sheetnames) == 4                                                     # super admin default = all branches


def test_export_permissions_and_audit(client, people, invoices):
    q = {"date": DAY}
    assert client.get(f"{API}/exports/daily-invoices", params=q).status_code == 401
    assert client.get(f"{API}/exports/daily-invoices", params=q, headers=people["til"]["h"]).status_code == 403      # ordinary user
    assert client.get(f"{API}/exports/daily-invoices", params={**q, "branch": "SBR"}, headers=people["til_admin"]["h"]).status_code == 403
    assert client.get(f"{API}/exports/daily-invoices", params={**q, "branch": "ALL"}, headers=people["til_admin"]["h"]).status_code == 403
    assert client.get(f"{API}/exports/daily-invoices", params={**q, "branch": "XXX"}, headers=people["super"]["h"]).status_code == 422
    assert client.get(f"{API}/exports/daily-invoices", params={"date": "not-a-date"}, headers=people["super"]["h"]).status_code == 422
    ok = client.get(f"{API}/exports/daily-invoices", params={**q, "branch": "TIL"}, headers=people["til_admin"]["h"])
    assert ok.status_code == 200 and ok.headers["content-disposition"] == f'attachment; filename="invoices-TIL-{DAY}.xlsx"'
    with engine.connect() as c:
        row = c.execute(text("select actor_name, metadata from audit_logs where action = 'report.export' order by created_at desc limit 1")).one()
    assert row[0] == "til_admin" and row[1]["date"] == DAY and row[1]["invoices"] == 2


# ---------------------------------------------------------------- stock
@pytest.fixture(scope="module", autouse=True)
def forget_stock_levels():
    """Invoices are limited by the stock a branch has set up, so these tests must not leave stock behind for the others."""
    yield
    with engine.begin() as c:
        c.execute(text("delete from stock_transfers"))
        c.execute(text("delete from branch_stock"))


def put_stock(client, h, items, branch=None):
    return client.put(f"{API}/stock", json={"items": items}, headers=h, params={"branch": branch} if branch else None)


def test_stock_remaining_is_opening_minus_invoiced_and_flags_low(client, people, prods):
    h, adm = people["bho"]["h"], people["super"]["h"]
    q = lambda: {i["name"]: i for i in client.get(f"{API}/stock", headers=h).json()["items"]}
    assert q()["Badges"]["status"] == "UNSET"                                                # nothing configured yet
    b0, c0 = q()["Badges"]["sold"], q()["Caps"]["sold"]                                      # other tests may already have sold some
    r = put_stock(client, adm, [{"product_id": prods["Badges"], "opening_qty": b0 + 100, "low_threshold": 20},
                                {"product_id": prods["Caps"], "opening_qty": c0 + 5, "low_threshold": 10}], branch="BHO")
    assert r.status_code == 200
    s = q()
    assert (s["Badges"]["remaining"], s["Badges"]["status"]) == (100, "OK") and s["Caps"]["status"] == "LOW"
    inv = make(client, people, prods, "bho", "Stock Co", 70)                                 # sells 70 Badges
    assert (q()["Badges"]["sold"], q()["Badges"]["remaining"], q()["Badges"]["status"]) == (b0 + 70, 30, "OK")
    make(client, people, prods, "bho", "Stock Co 2", 15)
    assert (q()["Badges"]["remaining"], q()["Badges"]["status"]) == (15, "LOW")             # 100 - 85 <= 20
    make(client, people, prods, "bho", "Draft", 50, action="draft")
    assert q()["Badges"]["remaining"] == 15                                                  # drafts do not reduce stock
    # a revision to a smaller quantity, and a cancellation, both give stock back automatically
    low = client.get(f"{API}/stock/low", headers=h).json()
    assert [i["name"] for i in low["items"]] == ["Caps", "Badges"]                           # emptiest first (5 left, then 15)
    assert low["configured"] == 2 and low["low_count"] == 2
    revised = {"company_name": "Stock Co", "invoice_date": DAY, "action": "submit", "edit_reason": "smaller order",
               "items": [{"product_id": prods["Badges"], "quantity": 10}]}
    assert client.put(f"{API}/invoices/{inv['id']}", json=revised, headers=people["super"]["h"]).status_code == 200
    assert q()["Badges"]["remaining"] == 75                                                  # 100 - 10 - 15
    client.post(f"{API}/invoices/{inv['id']}/cancel", json={"reason": "withdrawn"}, headers=people["super"]["h"])
    assert q()["Badges"]["remaining"] == 85                                                  # 100 - 15
    make(client, people, prods, "bho", "Big", 85)
    assert (q()["Badges"]["remaining"], q()["Badges"]["status"]) == (0, "OUT")


def test_stock_permissions_and_isolation(client, people, prods):
    body = [{"product_id": prods["Badges"], "opening_qty": 10, "low_threshold": 2}]
    assert client.get(f"{API}/stock").status_code == 401
    assert put_stock(client, people["til"]["h"], body).status_code == 403                       # ordinary user cannot change stock
    assert put_stock(client, people["til_admin"]["h"], body, branch="SBR").status_code == 403   # nor reach another branch
    assert client.get(f"{API}/stock", params={"branch": "SBR"}, headers=people["til_admin"]["h"]).status_code == 403
    assert client.get(f"{API}/stock", headers=people["super"]["h"]).status_code == 422          # super admin must pick a branch
    ok = put_stock(client, people["til_admin"]["h"], body)                                      # a branch admin sets their OWN branch
    assert ok.status_code == 200 and ok.json()["branch_code"] == "TIL"
    sbr = {i["name"]: i for i in client.get(f"{API}/stock", headers=people["sbr"]["h"]).json()["items"]}
    assert sbr["Badges"]["status"] == "UNSET"                                                   # TIL's numbers never leak to SBR
    for bad in ([{"product_id": prods["Badges"], "opening_qty": -1, "low_threshold": 2}],
                [{"product_id": prods["Badges"], "opening_qty": 5, "low_threshold": 2, "extra": 1}],
                [{"product_id": "00000000-0000-0000-0000-000000000000", "opening_qty": 5, "low_threshold": 2}]):
        assert put_stock(client, people["til_admin"]["h"], bad).status_code == 422
    with engine.connect() as c:                                                                 # the database agrees
        c.execute(text("select set_config('app.user_id', :u, true), set_config('app.role', 'BRANCH_USER', true), set_config('app.branch_id', "
                       "(select id::text from branches where code = 'SBR'), true)"), {"u": str(people["sbr"]["user_id"])})
        c.execute(text("set local role app_authenticated"))
        assert c.execute(text("select count(*) from branch_stock")).scalar() == 0
        assert c.execute(text("update branch_stock set opening_qty = 0")).rowcount == 0
    # stock edits are audited
    with engine.connect() as c:
        n = c.execute(text("select count(*) from audit_logs where entity_type = 'stock'")).scalar()
    assert n >= 2


# ---------------------------------------------------------------- transfers (central admin only)
def test_central_admin_sees_all_branches_and_transfers_stock(client, people, prods):
    adm = people["super"]["h"]
    badges = prods["Badges"]
    levels = lambda code: {i["name"]: i for i in client.get(f"{API}/stock", headers=adm, params={"branch": code}).json()["items"]}
    s0, h0 = levels("SBR")["Badges"]["sold"], levels("HAD")["Badges"]["sold"]
    assert put_stock(client, adm, [{"product_id": badges, "opening_qty": s0 + 50, "low_threshold": 5}], branch="SBR").status_code == 200
    assert put_stock(client, adm, [{"product_id": badges, "opening_qty": h0 + 3, "low_threshold": 5}], branch="HAD").status_code == 200

    ov = client.get(f"{API}/stock/overview", headers=adm).json()
    assert {b["code"] for b in ov["branches"]} >= {"SBR", "HAD"}
    row = next(i for i in ov["items"] if i["name"] == "Badges")
    assert row["branches"]["SBR"]["remaining"] == 50 and row["branches"]["HAD"]["remaining"] == 3

    body = {"product_id": badges, "from_branch": "SBR", "to_branch": "HAD", "quantity": 20, "note": "HAD running out"}
    r = client.post(f"{API}/stock/transfers", json=body, headers=adm)
    assert r.status_code == 201 and r.json()["from_branch"] and r.json()["quantity"] == 20
    assert (levels("SBR")["Badges"]["remaining"], levels("SBR")["Badges"]["transferred_out"]) == (30, 20)
    assert (levels("HAD")["Badges"]["remaining"], levels("HAD")["Badges"]["transferred_in"]) == (23, 20)
    assert client.get(f"{API}/stock/transfers", headers=adm).json()[0]["note"] == "HAD running out"

    # cannot send more than the sender has, to itself, or to a branch that doesn't exist
    assert client.post(f"{API}/stock/transfers", json={**body, "quantity": 31}, headers=adm).status_code == 422
    assert client.post(f"{API}/stock/transfers", json={**body, "to_branch": "SBR"}, headers=adm).status_code == 422
    assert client.post(f"{API}/stock/transfers", json={**body, "to_branch": "XXX"}, headers=adm).status_code == 422
    assert client.post(f"{API}/stock/transfers", json={**body, "quantity": 0}, headers=adm).status_code == 422
    # a branch that had nothing set can receive stock
    r = client.post(f"{API}/stock/transfers", json={**body, "to_branch": "AHL", "quantity": 5}, headers=adm)
    assert r.status_code == 201 and levels("AHL")["Badges"]["remaining"] == 5

    # nobody but the central admin can see the overview or move stock
    for who in ("til", "til_admin"):
        h = people[who]["h"]
        assert client.get(f"{API}/stock/overview", headers=h).status_code == 403
        assert client.get(f"{API}/stock/transfers", headers=h).status_code == 403
        assert client.post(f"{API}/stock/transfers", json=body, headers=h).status_code == 403
    assert client.post(f"{API}/stock/transfers", json=body).status_code == 401
    # a receiving branch sees its new remaining stock; the database refuses writes from a branch admin
    til_admin = people["til_admin"]
    with engine.connect() as c:
        c.execute(text("select set_config('app.user_id', :u, true), set_config('app.role', 'BRANCH_ADMIN', true), set_config('app.branch_id', "
                       "(select id::text from branches where code = 'TIL'), true)"), {"u": str(til_admin["user_id"])})
        c.execute(text("set local role app_authenticated"))
        with pytest.raises(Exception):
            c.execute(text("insert into stock_transfers (event_id, product_id, from_branch_id, to_branch_id, quantity) "
                           "values ((select id from events limit 1), :p, (select id from branches where code='TIL'), "
                           "(select id from branches where code='SBR'), 1)"), {"p": badges})
    with engine.connect() as c:
        n = c.execute(text("select count(*) from audit_logs where action = 'stock.transfer'")).scalar()
    assert n >= 2


# ---------------------------------------------------------------- stock limits, alerts, Razorpay mode
def test_a_submitted_invoice_cannot_take_more_than_the_branch_has(client, people, prods):
    adm, h = people["super"]["h"], people["til"]["h"]
    items = {i["name"]: i for i in client.get(f"{API}/stock", headers=adm, params={"branch": "TIL"}).json()["items"]}
    sold = items["Caps"]["sold"] + items["Caps"]["transferred_out"] - items["Caps"]["transferred_in"]
    assert put_stock(client, adm, [{"product_id": prods["Caps"], "opening_qty": sold + 10, "low_threshold": 3}], branch="TIL").status_code == 200

    def body(qty, action="submit", **extra):
        return {"company_name": "Limit Co", "invoice_date": DAY, "action": action,
                "items": [{"product_id": prods["Caps"], "quantity": qty}], **extra}

    over = client.post(f"{API}/invoices", headers=h, json=body(11))
    assert over.status_code == 422 and "Only 10 of Caps" in over.json()["detail"]
    assert client.post(f"{API}/invoices", headers=h, json=body(50, action="draft")).status_code == 201     # a draft is only a note
    ok = client.post(f"{API}/invoices", headers=h, json=body(8)).json()
    assert client.post(f"{API}/invoices", headers=h, json=body(3)).status_code == 422                        # only 2 left now
    # a revision may keep its own quantity but not exceed what is left plus that quantity
    rev = lambda q: client.put(f"{API}/invoices/{ok['id']}", headers=people["super"]["h"], json=body(q, edit_reason="change"))
    assert rev(10).status_code == 200                                                                          # 2 left + its own 8
    assert rev(11).status_code == 422
    # a branch with no stock set up is not limited
    assert client.post(f"{API}/invoices", headers=people["sbr"]["h"], json=body(9999)).status_code == 201


def test_low_stock_alerts_for_the_central_admin_and_each_branch(client, people, prods):
    adm = people["super"]["h"]
    items = {i["name"]: i for i in client.get(f"{API}/stock", headers=adm, params={"branch": "BHO"}).json()["items"]}
    base = items["Flags - Handy"]["sold"] + items["Flags - Handy"]["transferred_out"] - items["Flags - Handy"]["transferred_in"]
    put_stock(client, adm, [{"product_id": prods["Flags - Handy"], "opening_qty": base + 3, "low_threshold": 5}], branch="BHO")   # 3 left, low at 5
    a = client.get(f"{API}/stock/alerts", headers=adm).json()
    mine = [i for i in a["items"] if i["name"] == "Flags - Handy" and i["branch_code"] == "BHO"]
    assert mine and mine[0]["status"] == "LOW" and mine[0]["remaining"] == 3 and a["count"] >= 1 and a["low"] >= 1
    assert all(i["branch_code"] == "BHO" for i in client.get(f"{API}/stock/alerts", headers=people["bho"]["h"]).json()["items"])
    assert not client.get(f"{API}/stock/alerts", headers=people["sbr"]["h"]).json()["items"] or all(
        i["branch_code"] == "SBR" for i in client.get(f"{API}/stock/alerts", headers=people["sbr"]["h"]).json()["items"])
    assert client.get(f"{API}/stock/alerts").status_code == 401


def test_razorpay_is_a_payment_mode(client, people, prods):
    r = client.post(f"{API}/invoices", headers=people["sbr"]["h"], json={
        "company_name": "Rzp Co", "invoice_date": DAY, "action": "submit",
        "items": [{"product_id": prods["Badges"], "quantity": 1}],
        "payments": [{"mode": "RAZORPAY", "amount": "6", "reference": "pay_Abc123"}]})
    assert r.status_code == 201 and r.json()["payments"][0]["mode"] == "RAZORPAY" and r.json()["payment_status"] == "PAID"
