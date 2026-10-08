"""Bulk download of invoices for one day or a date range: one Excel workbook, or a ZIP of the PDFs.

One workbook: a Summary tab, then one tab per invoice laid out like the MCCIA invoice sheet.
Tab name = branch code + invoice sequence + company, e.g. "TIL-000007 Sahyadri Precision".
  - Branch admin: their own branch only.
  - Central admin: one branch at a time (five separate files) or ALL branches combined, grouped by branch.
Visibility is enforced twice, as everywhere else: services.list_invoices scopes by the verified principal,
and the database's row level security hides other branches' rows.
"""

import io
import re
import zipfile
from datetime import date
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import admin, repositories as repo, services
from .auth import Principal, current_principal
from .db import get_session
from .models import Invoice, Product

router = APIRouter(prefix="/api/v1/exports")
BLUE = PatternFill("solid", fgColor="C5D9F1")
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(name="Calibri", bold=True)
MONEY = '#,##0.00;-#,##0.00;"-"'
MODE = {"CASH": "Cash", "RAZORPAY": "Razorpay", "UPI": "UPI", "CARD": "Card", "NET_BANKING": "Net banking", "OTHER": "Other"}
BAD_TAB_CHARS = re.compile(r"[\[\]:*?/\\]")


def tab_name(inv: Invoice, taken: set[str]) -> str:
    """`TIL-000007 Company`, max 31 characters, no characters Excel forbids, unique within the workbook."""
    _, code, seq = inv.invoice_number.split("-")
    base = BAD_TAB_CHARS.sub("", f"{code}-{seq} {inv.company_name}").strip()[:31] or f"{code}-{seq}"
    name, n = base, 2
    while name.lower() in taken:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(name.lower())
    return name


def payment_text(inv: Invoice) -> str:
    """Same wording as the invoice sheet / PDF."""
    parts = [f"{MODE[p.mode]} Rs. {p.amount:,.2f}" + (f" ({p.reference})" if p.reference else "") for p in inv.payments]
    lines = [" + ".join(parts)] if parts else []
    balance = inv.grand_total - inv.amount_paid
    if parts and balance > 0:
        lines.append(f"Balance due Rs. {balance:,.2f}")
    if inv.payment_details.strip():
        lines.append(inv.payment_details.strip())
    return "  |  ".join(lines)


def invoice_sheet(wb: Workbook, inv: Invoice, products: list[Product], header: dict, taken: set[str]) -> None:
    ws = wb.create_sheet(tab_name(inv, taken))
    disc = inv.discount_percent > 0
    cols = ["Sr.", "Particulars", "HSN Code", "Rate", "Qty."] + (["Rate after Discount"] if disc else []) + [
        "Basic Amount", "CGST Rate", "CGST AMT", "SGST Rate", "SGST AMT", "Total Amount"]
    n = len(cols)
    last = chr(64 + n)
    for i, w in enumerate([6, 38, 11, 9, 7] + ([11] if disc else []) + [12, 8, 10, 8, 10, 13], 1):
        ws.column_dimensions[chr(64 + i)].width = w

    def put(row, text_, bold=False, fill=None, merge_to=None, align="left", border=True):
        ws.cell(row, 1, text_)
        c = ws.cell(row, 1)
        c.font = BOLD if bold else Font(name="Calibri")
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            c.fill = fill
        if merge_to:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=merge_to)
        if border:
            for col in range(1, (merge_to or 1) + 1):
                ws.cell(row, col).border = BOX

    put(1, header["title"], True, merge_to=n, align="center")
    put(2, header["name"], True, merge_to=n)
    for i, line in enumerate(header["address_lines"][:2]):
        put(3 + i, line.strip(), merge_to=n)
    put(5, f"GSTIN - {header['gstin']}   PAN - {header['pan']}", True, merge_to=n)
    put(7, f"Company Name  : {inv.company_name}", True, merge_to=n - 3)
    ws.cell(7, n - 2, "Proforma Invoice No.").font = BOLD
    ws.cell(7, n, inv.invoice_number).font = BOLD
    put(8, f"Address : {inv.address}", True, merge_to=n - 3)
    ws.cell(8, n - 2, "Proforma Invoice Date").font = BOLD
    ws.cell(8, n, inv.invoice_date.strftime("%d/%b/%Y")).font = BOLD
    put(9, f"GSTIN  {inv.gstin}      Email ID  {inv.email}", True, merge_to=n - 3)
    put(10, f"Contact Person & Cell No : {inv.contact_person} {inv.contact_phone}".rstrip(), True, merge_to=n - 3)

    r0 = 12
    for i, name in enumerate(cols, 1):
        c = ws.cell(r0, i, name)
        c.font, c.fill, c.border = BOLD, BLUE, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 32

    bought = {i.product_id: i for i in inv.items if i.product_id}
    rows = [(p, bought.get(p.id)) for p in products]
    listed = {p.id for p in products}
    extra = [(None, i) for i in inv.items if not i.product_id or i.product_id not in listed]  # never drop a line
    r = r0
    for sr, (prod, item) in enumerate(rows + extra, 1):
        r += 1
        name = item.particulars if item else " ".join(prod.name.split())
        hsn = item.hsn_code if item else prod.hsn_code
        rate = item.rate if item else prod.current_rate
        gst = (item.cgst_rate if item else prod.cgst_rate)
        vals = [sr, name, hsn, rate, item.quantity if item else None]
        if disc:
            vals.append(item.rate_after_discount if item else rate - rate * inv.discount_percent / 100)
        vals += [item.basic_amount if item else 0, gst / 100, item.cgst_amount if item else 0, gst / 100,
                 item.sgst_amount if item else 0, item.total_amount if item else 0]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, float(v) if hasattr(v, "quantize") else v)
            c.border = BOX
            c.font = Font(name="Calibri", bold=bool(item and i == 5))
            if cols[i - 1] in ("Rate", "Basic Amount", "CGST AMT", "SGST AMT", "Total Amount", "Rate after Discount"):
                c.number_format = MONEY
            if cols[i - 1] in ("CGST Rate", "SGST Rate"):
                c.number_format = "0.0%"
            if i == 5:
                c.alignment = Alignment(horizontal="center")

    r += 1  # totals row (blue, like the invoice)
    col_of = {name: i for i, name in enumerate(cols, 1)}
    total_qty = sum(i.quantity for i in inv.items)
    for i in range(1, n + 1):
        c = ws.cell(r, i)
        c.fill, c.border, c.font = BLUE, BOX, BOLD
        c.number_format = MONEY
    ws.cell(r, 2, "Total …").alignment = Alignment(horizontal="center")
    ws.cell(r, col_of["Qty."], total_qty)
    for key, v in (("Basic Amount", inv.subtotal), ("CGST AMT", inv.cgst_total), ("SGST AMT", inv.sgst_total),
                   ("Total Amount", inv.grand_total - inv.rounding_adjustment)):
        ws.cell(r, col_of[key], float(v))
    r += 1  # rounded-off row (blue label + value)
    ws.cell(r, n - 3, "Rounded off Amount..").font = BOLD
    ws.merge_cells(start_row=r, start_column=n - 3, end_row=r, end_column=n - 1)
    ws.cell(r, n, float(inv.grand_total)).font = BOLD
    ws.cell(r, n).number_format = MONEY
    for i in range(n - 3, n + 1):
        ws.cell(r, i).fill, ws.cell(r, i).border = BLUE, BOX
    ws.cell(r, n - 3).alignment = Alignment(horizontal="center")
    r += 1
    put(r, f"Amount in Words : {inv.amount_in_words}", merge_to=n - 3)
    ws.cell(r, n - 2, header["for_org"]).font = BOLD
    ws.merge_cells(start_row=r, start_column=n - 2, end_row=r, end_column=n)
    r += 1
    put(r, f"Payment Details : {payment_text(inv)}", fill=BLUE, merge_to=n - 3)
    ws.cell(r, n - 2, header["signatory"]).font = BOLD
    ws.merge_cells(start_row=r, start_column=n - 2, end_row=r, end_column=n)
    ws.row_dimensions[r].height = 36
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "portrait", 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def summary_sheet(wb: Workbook, invoices: list[Invoice], day: date, scope: str, codes: dict, day_to: date | None = None) -> None:
    ws = wb.active
    ws.title = "Summary"
    span = day.strftime('%d/%b/%Y') if not day_to or day_to == day else f"{day.strftime('%d/%b/%Y')} to {day_to.strftime('%d/%b/%Y')}"
    ws["A1"] = f"Proforma Invoices - {span} - {scope}"
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    head = ["Branch", "Invoice No.", "Company", "Status", "Payment", "Items qty", "Amount (Rs.)"]
    for i, h in enumerate(head, 1):
        c = ws.cell(3, i, h)
        c.font, c.fill, c.border = BOLD, BLUE, BOX
    for i, w in enumerate([10, 20, 38, 14, 12, 10, 16], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    r, grand, grand_n = 3, 0.0, 0
    by_branch: dict[str, list[Invoice]] = {}
    for inv in invoices:
        by_branch.setdefault(codes[inv.branch_id], []).append(inv)
    for code in sorted(by_branch):
        sub = 0.0
        for inv in sorted(by_branch[code], key=lambda i: i.invoice_number):
            r += 1
            status = {"PAID": "Paid", "PARTIAL": "Part paid", "UNPAID": "Unpaid"}[inv.payment_status]
            for i, v in enumerate([code, inv.invoice_number, inv.company_name, inv.status.title(), status,
                                   sum(i.quantity for i in inv.items), float(inv.grand_total)], 1):
                c = ws.cell(r, i, v)
                c.border = BOX
                if i == 7:
                    c.number_format = MONEY
            sub += float(inv.grand_total)
        grand += sub
        grand_n += len(by_branch[code])
        r += 1  # per-branch subtotal: the "bifurcation" in the combined file
        ws.cell(r, 1, f"{code} total").font = BOLD
        ws.cell(r, 2, f"{len(by_branch[code])} invoice(s)").font = BOLD
        c = ws.cell(r, 7, sub)
        c.font, c.number_format = BOLD, MONEY
        for i in range(1, 8):
            ws.cell(r, i).fill, ws.cell(r, i).border = BLUE, BOX
    r += 2
    ws.cell(r, 1, "All branches" if len(by_branch) != 1 else "Total").font = BOLD
    ws.cell(r, 2, f"{grand_n} invoice(s)").font = BOLD
    c = ws.cell(r, 7, grand)
    c.font, c.number_format = BOLD, MONEY
    ws.freeze_panes = "A4"


MAX_RANGE_DAYS = 92
MAX_ZIP_INVOICES = 60  # PDFs are built on demand if missing, so a ZIP is kept small enough to finish quickly


def _range_and_scope(s: Session, p: Principal, day: date | None, date_from: date | None, date_to: date | None, branch: str | None):
    """Who may download what, which dates, and the branch scope. A single `date`, or `date_from` + `date_to` (inclusive)."""
    if not p.is_admin:
        raise HTTPException(403, "Only an admin can download invoices in bulk")
    d_from, d_to = (day, day) if day else (date_from, date_to)
    if d_from is None or d_to is None:
        raise HTTPException(422, "Choose the dates (from and to)")
    if d_to < d_from:
        raise HTTPException(422, "The 'to' date is before the 'from' date")
    if (d_to - d_from).days > MAX_RANGE_DAYS:
        raise HTTPException(422, f"Please choose a range of at most {MAX_RANGE_DAYS} days")
    if p.is_super:
        code = (branch or "ALL").upper()
        if code != "ALL" and repo.get_branch_by_code(s, code) is None:
            raise HTTPException(422, "Unknown branch")
    else:
        if branch is not None and branch.upper() != _own_code(s, p):
            raise HTTPException(403, "You can download your own branch only")
        code = _own_code(s, p)
    rows = services.list_invoices(s, p, None if code == "ALL" else (code if p.is_super else None),
                                  date_from=d_from, date_to=d_to, limit=500)
    rows = [i for i in rows if i.status in admin.COUNTED]  # drafts and cancelled invoices are not "issued"
    return d_from, d_to, code, rows


def _stamp(d_from: date, d_to: date) -> str:
    return d_from.isoformat() if d_from == d_to else f"{d_from.isoformat()}_to_{d_to.isoformat()}"


@router.get("/daily-invoices")
def daily_invoices(
    day: Annotated[date | None, Query(alias="date")] = None,
    date_from: date | None = None,
    date_to: date | None = None,
    branch: Annotated[str | None, Query(max_length=3)] = None,
    s: Session = Depends(get_session),
    p: Principal = Depends(current_principal),
):
    """One Excel workbook (a Summary tab, then one tab per invoice) for one day or a date range."""
    d_from, d_to, code, rows = _range_and_scope(s, p, day, date_from, date_to, branch)
    event = repo.current_event(s)
    products = repo.list_products(s, event.id) if event else []
    header = admin.read_header(s)
    codes = {b.id: b.code for b in repo.list_branches(s)}

    wb = Workbook()
    summary_sheet(wb, rows, d_from, "all branches" if code == "ALL" else f"branch {code}", codes, d_to)
    taken = {"summary"}
    for inv in sorted(rows, key=lambda i: i.invoice_number.split("-")[1:]):
        invoice_sheet(wb, inv, products, header, taken)

    services.log_event(s, p, "report.export", entity_type="report", branch_id=None if code == "ALL" else rows[0].branch_id if rows else p.branch_id,
                       kind="daily-invoices", date=d_from.isoformat(), date_to=d_to.isoformat(), branch=code, invoices=len(rows))
    s.commit()
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="invoices-{code}-{_stamp(d_from, d_to)}.xlsx"',
                             "Cache-Control": "no-store"})


@router.get("/invoice-pdfs")
def invoice_pdfs(
    day: Annotated[date | None, Query(alias="date")] = None,
    date_from: date | None = None,
    date_to: date | None = None,
    branch: Annotated[str | None, Query(max_length=3)] = None,
    s: Session = Depends(get_session),
    p: Principal = Depends(current_principal),
):
    """A ZIP with each invoice's current PDF for one day or a date range (same scope rules as the Excel download)."""
    from . import documents  # local import: documents imports this package's siblings

    d_from, d_to, code, rows = _range_and_scope(s, p, day, date_from, date_to, branch)
    if not rows:
        raise HTTPException(404, "No issued invoices in those dates")
    if len(rows) > MAX_ZIP_INVOICES:
        raise HTTPException(422, f"{len(rows)} invoices match - a ZIP holds at most {MAX_ZIP_INVOICES}. Please choose fewer days (or one branch).")
    codes = {b.id: b.code for b in repo.list_branches(s)}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for inv in sorted(rows, key=lambda i: i.invoice_number):
            try:
                data, name = documents.read_document(s, p, inv.id, None)
            except (services.Unavailable, services.Conflict, services.NotFound) as e:
                raise HTTPException(503, f"Could not prepare {inv.invoice_number}: {e}") from None
            z.writestr(f"{codes[inv.branch_id]}/{name}" if code == "ALL" else name, data)
    services.log_event(s, p, "report.export", entity_type="report", branch_id=None if code == "ALL" else rows[0].branch_id,
                       kind="invoice-pdfs", date=d_from.isoformat(), date_to=d_to.isoformat(), branch=code, invoices=len(rows))
    s.commit()
    return Response(buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="invoice-pdfs-{code}-{_stamp(d_from, d_to)}.zip"',
                             "Cache-Control": "no-store"})


def _own_code(s: Session, p: Principal) -> str:
    return s.get(repo.Branch, p.branch_id).code
